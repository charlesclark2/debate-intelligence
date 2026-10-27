"""Cut cards as a Verbatim .docx: the DOCX side of ExportService (`v1-e06-t04`).

:func:`render_cards_docx` writes the cards of one :class:`~debate_core.evidence.manifest.CardManifest`
in a :class:`~debate_core.evidence.format_profile.ResolvedFormatProfile`, with python-docx. Each card
is its tag, its cite and its evidence, under an optional pocket, hat and block written once at the top.

## Where every style comes from

Nothing here names a style or a size. The document starts from a package built from the base style
profile (`v1-e31-t02`): its `word/styles.xml` is the profile's document defaults and its writer
definitions, in the profile's order and in the same form as the lossless writer's base template
(`v1-e33-t02`), so both writers emit one style table for one profile. Then:

* pocket, hat, block and tag are the profile's paragraph styles for those units (`Heading1` to
  `Heading4`), whose definitions carry the outline levels Word's navigation pane reads;
* the cite and the evidence take the profile's paragraph style for them when the profile writes a
  definition for it, and the default paragraph style otherwise, as Verbatim does: a cite is
  recognised by its short cite's character style, evidence by its place;
* the short cite carries the profile's cite character style when the format profile says it is
  bold; the rest of the cite is at the format profile's full-cite size;
* evidence that a span underlines carries the underline character style, evidence a span highlights
  carries the format profile's highlight colour (none in `plain`), read text is at the format
  profile's base size and unread text at its condensed size.

## The evidence is `Card.evidence_text` and nothing else

The evidence paragraph is laid out by :func:`~debate_core.evidence.run_segmentation.lay_out_evidence`
from offsets alone, and every evidence run's text is a slice of `evidence_text`. Two things are added
to it, each as a run of its own, never underlined or highlighted (they are not evidence and no span
covers them), at the base size:

* the format profile's omission marker at each omitted range (ADR-0018);
* each interpolation as :attr:`~debate_core.domain.Interpolation.rendered`, the only code that adds
  brackets (`v1-e03-t05`).

Neither is ever part of `evidence_text`. Text a .docx cannot carry as written (a carriage return,
which python-docx writes as a line break and reads back as a line feed, or a character XML 1.0 does
not allow) is refused with :class:`~debate_core.evidence.card_rendering.RenderRefused`, never altered.

## The same cards give the same bytes, and no one's name

The core properties name the tool as author and last editor, nothing about the student, and take the
manifest's `generated_at` as both their times; no other document property part is written. Every
member of the package is stamped with that time too, so one manifest and one profile always render
the same bytes.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Sequence
from datetime import UTC, datetime
from typing import Final
from xml.sax.saxutils import quoteattr

import docx
from docx.document import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.opc.constants import CONTENT_TYPE, RELATIONSHIP_TYPE
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Pt, RGBColor
from docx.styles.style import CharacterStyle, ParagraphStyle
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from debate_core.domain import SpanStyle
from debate_core.domain.style_profile import (
    RunEmphasis,
    StructuralUnit,
    StyleProfile,
    WriterDocumentDefaults,
    WriterStyleDefinition,
)
from debate_core.evidence.card_rendering import (
    CardToRender,
    DocumentHeadings,
    RenderRefused,
    banner_text,
    cite_marks,
)
from debate_core.evidence.docx_manifest_part import (
    MANIFEST_PART_NAME,
    MANIFEST_PROPERTIES_PART_NAME,
    EmbeddedManifestError,
    manifest_part_xml,
    manifest_properties_xml,
)
from debate_core.evidence.format_profile import ResolvedFormatProfile
from debate_core.evidence.manifest import CardManifest, serialize_manifest
from debate_core.evidence.run_segmentation import (
    EvidenceRun,
    InterpolationMark,
    OmissionMark,
    lay_out_evidence,
)

__all__ = ["DOCUMENT_AUTHOR", "render_cards_docx"]

DOCUMENT_AUTHOR: Final = "debate-research"
"""The author and last editor every rendered document names: the tool, never a student."""

_W_NAMESPACE: Final = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_DECLARATION: Final = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'

#: Word's compatibility mode for a current document, as the lossless writer's base sets it.
_WORD_COMPATIBILITY: Final = (
    '<w:compat><w:compatSetting w:name="compatibilityMode" '
    'w:uri="http://schemas.microsoft.com/office/word" w:val="15"/></w:compat>'
)

_CONTENT_TYPES: Final = (
    _DECLARATION + '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
    '<Default Extension="xml" ContentType="application/xml"/>'
    '<Override PartName="/word/document.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
    '<Override PartName="/word/styles.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.styles+xml"/>'
    '<Override PartName="/word/settings.xml" '
    'ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.settings+xml"/>'
    '<Override PartName="/docProps/core.xml" '
    'ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>'
    "</Types>"
)

_PACKAGE_RELATIONSHIPS: Final = (
    _DECLARATION + '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
    'Target="word/document.xml"/>'
    '<Relationship Id="rId2" '
    'Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" '
    'Target="docProps/core.xml"/>'
    "</Relationships>"
)

_DOCUMENT_RELATIONSHIPS: Final = (
    _DECLARATION + '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
    '<Relationship Id="rId1" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>'
    '<Relationship Id="rId2" '
    'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/settings" '
    'Target="settings.xml"/>'
    "</Relationships>"
)

_EMPTY_CORE_PROPERTIES: Final = (
    _DECLARATION + "<cp:coreProperties "
    'xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
    'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"/>'
)

#: The paragraph units a reader navigates by: a fallback to the default style would lose them.
_HEADING_UNITS: Final = frozenset(
    {StructuralUnit.POCKET, StructuralUnit.HAT, StructuralUnit.BLOCK, StructuralUnit.TAG}
)

#: The span style each character style draws; highlighting is a colour, not a style.
_CHARACTER_STYLE_OF_SPAN: Final = {SpanStyle.UNDERLINE: RunEmphasis.UNDERLINE}


def render_cards_docx(
    manifest: CardManifest,
    cards: Sequence[CardToRender],
    profile: ResolvedFormatProfile,
    *,
    headings: DocumentHeadings | None = None,
) -> bytes:
    """The manifest's cards as a .docx, in its order, in `profile`'s formatting.

    `cards` gives each manifest card its cite text and, unless it is VERIFIED, its reasons; they must
    be the manifest's cards exactly, in its order, so the document is what its manifest describes.
    Raises :class:`~debate_core.evidence.card_rendering.RenderRefused` before writing anything.
    """
    if tuple(item.card for item in cards) != manifest.cards:
        raise RenderRefused(
            "the cards to render must be the manifest's cards, in its order: the manifest describes the "
            "document it is embedded in"
        )
    for item in cards:
        _check_card_writable(item)
    manifest_bytes = serialize_manifest(manifest)
    try:
        manifest_part = manifest_part_xml(manifest_bytes)
    except EmbeddedManifestError as error:
        raise RenderRefused(str(error)) from error
    writer = _CardDocumentWriter(profile)
    writer.write_headings(headings or DocumentHeadings())
    for item in cards:
        writer.write_card(item)
    writer.embed_manifest(manifest_part, manifest_properties_xml(manifest_bytes))
    return writer.save(manifest.generated_at)


class _CardDocumentWriter:
    """One document being written: python-docx over a base package built from the style profile."""

    def __init__(self, profile: ResolvedFormatProfile) -> None:
        self._profile = profile
        self._document: Document = docx.Document(io.BytesIO(_base_package(profile.style_profile)))
        styles = list(self._document.styles)
        self._paragraph_styles = {s.style_id: s for s in styles if isinstance(s, ParagraphStyle)}
        self._character_styles = {
            s.style_id: s
            for s in styles
            if isinstance(s, CharacterStyle) and not isinstance(s, ParagraphStyle)
        }
        defaults = [
            d.style_id
            for d in profile.style_profile.writer_style_definitions
            if d.is_default and d.style_type == "paragraph"
        ]
        if len(defaults) != 1:
            raise RenderRefused(
                f"style profile {profile.style_profile.name} must write exactly one default paragraph style"
            )
        self._default_paragraph_style = self._paragraph_styles[defaults[0]]

    # -- styles -------------------------------------------------------------------------------

    def _paragraph(self, unit: StructuralUnit) -> Paragraph:
        style_id = self._profile.paragraph_style_id(unit)
        style = self._paragraph_styles.get(style_id)
        if style is None:
            if unit in _HEADING_UNITS:
                raise RenderRefused(
                    f"style profile {self._profile.style_profile.name} writes no definition for the "
                    f"{unit.value} style {style_id}, so the document would have no {unit.value} headings"
                )
            style = self._default_paragraph_style
        return self._document.add_paragraph(style=style)

    def _character_style(self, emphasis: RunEmphasis) -> CharacterStyle:
        style_id = self._profile.character_style_id(emphasis)
        style = self._character_styles.get(style_id)
        if style is None:
            raise RenderRefused(
                f"style profile {self._profile.style_profile.name} writes no definition for the "
                f"{emphasis.value} character style {style_id}"
            )
        return style

    # -- content ------------------------------------------------------------------------------

    def write_headings(self, headings: DocumentHeadings) -> None:
        for unit, text in (
            (StructuralUnit.POCKET, headings.pocket),
            (StructuralUnit.HAT, headings.hat),
            (StructuralUnit.BLOCK, headings.block),
        ):
            if text is not None:
                _check_writable(text, f"the {unit.value.lower()} heading")
                self._paragraph(unit).add_run(text)

    def write_card(self, item: CardToRender) -> None:
        card = item.card
        layout = self._profile.profile

        self._paragraph(StructuralUnit.TAG).add_run(card.tag)

        banner = banner_text(item)
        if banner is not None:
            run = self._paragraph(StructuralUnit.OTHER).add_run(banner)
            run.font.color.rgb = RGBColor.from_string(layout.unverified_banner_color)

        cite = self._paragraph(StructuralUnit.CITE)
        short = cite.add_run(item.cite.short_cite)
        if layout.cite_block.short_cite_bold:
            short.style = self._character_style(RunEmphasis.CITE)
        else:
            _size(short, layout.cite_block.full_cite_half_points)
        if item.cite.cite_line:
            _size(cite.add_run(item.cite.cite_line), layout.cite_block.full_cite_half_points)
        marks = cite_marks(card.citation)
        if marks:
            _size(cite.add_run(" " + " ".join(marks)), layout.cite_block.full_cite_half_points)

        if card.evidence_text:
            self._write_evidence(
                card.evidence_text, lay_out_evidence(card), self._paragraph(StructuralUnit.EVIDENCE)
            )

    def _write_evidence(
        self,
        evidence_text: str,
        pieces: Sequence[EvidenceRun | OmissionMark | InterpolationMark],
        paragraph: Paragraph,
    ) -> None:
        layout = self._profile.profile
        for piece in pieces:
            match piece:
                case EvidenceRun(start=start, end=end, styles=styles):
                    run = paragraph.add_run(evidence_text[start:end])
                    for span_style, emphasis in _CHARACTER_STYLE_OF_SPAN.items():
                        if span_style in styles:
                            run.style = self._character_style(emphasis)
                    if SpanStyle.HIGHLIGHT in styles and layout.highlight_color is not None:
                        run.font.highlight_color = WD_COLOR_INDEX.from_xml(layout.highlight_color)
                    size = layout.evidence_base_half_points
                    if piece.condensed and layout.condensed_half_points is not None:
                        size = layout.condensed_half_points
                    _size(run, size)
                case OmissionMark():
                    _size(paragraph.add_run(layout.omission_marker), layout.evidence_base_half_points)
                case InterpolationMark(interpolation=interpolation):
                    _size(paragraph.add_run(interpolation.rendered), layout.evidence_base_half_points)

    # -- the package --------------------------------------------------------------------------

    def embed_manifest(self, part_xml: bytes, properties_xml: bytes) -> None:
        """Add the manifest as a custom XML part related from the document, with its properties part."""
        package = self._document.part.package
        item = Part(PackURI(MANIFEST_PART_NAME), CONTENT_TYPE.XML, part_xml, package)
        properties = Part(
            PackURI(MANIFEST_PROPERTIES_PART_NAME),
            CONTENT_TYPE.OFC_CUSTOM_XML_PROPERTIES,
            properties_xml,
            package,
        )
        item.relate_to(properties, RELATIONSHIP_TYPE.CUSTOM_XML_PROPS)
        self._document.part.relate_to(item, RELATIONSHIP_TYPE.CUSTOM_XML)

    def save(self, generated_at: datetime) -> bytes:
        moment = generated_at.astimezone(UTC)
        properties = self._document.core_properties
        properties.author = DOCUMENT_AUTHOR
        properties.last_modified_by = DOCUMENT_AUTHOR
        properties.created = moment
        properties.modified = moment
        properties.revision = 1
        buffer = io.BytesIO()
        self._document.save(buffer)
        return _stamped(buffer.getvalue(), moment)


def _size(run: Run, half_points: int) -> None:
    run.font.size = Pt(half_points / 2)


def _check_card_writable(item: CardToRender) -> None:
    card = item.card
    _check_writable(card.tag, f"tag of card {card.card_id}")
    _check_writable(item.cite.full_cite, f"cite of card {card.card_id}")
    _check_writable(card.evidence_text, f"evidence of card {card.card_id}")
    for interpolation in card.interpolations:
        _check_writable(interpolation.text, f"an interpolation of card {card.card_id}")


def _check_writable(text: str, what: str) -> None:
    """Refuse text that a .docx cannot carry exactly as written, rather than change it."""
    for offset, character in enumerate(text):
        code = ord(character)
        if character in "\t\n" or 0x20 <= code <= 0xD7FF or 0xE000 <= code <= 0xFFFD or code >= 0x10000:
            continue
        raise RenderRefused(
            f"{what} holds U+{code:04X} at offset {offset}, which a .docx cannot carry as written"
        )


# --------------------------------------------------------------------------------------------
# The base package: the style profile's styles, and nothing python-docx's own template would add
# --------------------------------------------------------------------------------------------


def _element(name: str, attributes: dict[str, str] | None = None, children: str = "") -> str:
    attributes_text = "".join(f" w:{key}={quoteattr(value)}" for key, value in (attributes or {}).items())
    return (
        f"<w:{name}{attributes_text}>{children}</w:{name}>" if children else f"<w:{name}{attributes_text}/>"
    )


def _run_properties(
    *, font: str | None, bold: bool | None, color: str | None, half_points: int | None, underline: str | None
) -> str:
    """`w:rPr` in the schema's element order: rFonts, b, color, sz, szCs, u; nothing when all are unset."""
    children: list[str] = []
    if font is not None:
        children.append(_element("rFonts", {"ascii": font, "hAnsi": font, "cs": font}))
    if bold is not None:
        children.append(_element("b") if bold else _element("b", {"val": "0"}))
    if color is not None:
        children.append(_element("color", {"val": color}))
    if half_points is not None:
        children.append(_element("sz", {"val": str(half_points)}))
        children.append(_element("szCs", {"val": str(half_points)}))
    if underline is not None:
        children.append(_element("u", {"val": underline}))
    return _element("rPr", children="".join(children)) if children else ""


def _style(definition: WriterStyleDefinition) -> str:
    """One `w:style`, its children in the schema's order: name, basedOn, next, qFormat, pPr, rPr."""
    attributes = {"type": definition.style_type}
    if definition.is_default:
        attributes["default"] = "1"
    attributes["styleId"] = definition.style_id
    children = [_element("name", {"val": definition.style_name})]
    if definition.based_on is not None:
        children.append(_element("basedOn", {"val": definition.based_on}))
    if definition.next_style_id is not None:
        children.append(_element("next", {"val": definition.next_style_id}))
    children.append(_element("qFormat"))
    if definition.style_type == "paragraph":
        paragraph: list[str] = []
        if definition.keep_with_next:
            paragraph.append(_element("keepNext"))
        if definition.page_break_before:
            paragraph.append(_element("pageBreakBefore"))
        if definition.outline_level is not None:
            paragraph.append(_element("outlineLvl", {"val": str(definition.outline_level)}))
        if paragraph:
            children.append(_element("pPr", children="".join(paragraph)))
    children.append(
        _run_properties(
            font=definition.font,
            bold=definition.bold,
            color=definition.color,
            half_points=definition.half_points,
            underline=definition.underline,
        )
    )
    return _element("style", attributes, "".join(children))


def _document_defaults(defaults: WriterDocumentDefaults) -> str:
    run = _run_properties(
        font=defaults.font, bold=None, color=None, half_points=defaults.half_points, underline=None
    )
    return _element("docDefaults", children=_element("rPrDefault", children=run) + _element("pPrDefault"))


def _styles_xml(profile: StyleProfile) -> str:
    """`word/styles.xml`: the document defaults, then every writer definition in the profile's order."""
    styles = "".join(_style(definition) for definition in profile.writer_style_definitions)
    return (
        _DECLARATION
        + f'<w:styles xmlns:w="{_W_NAMESPACE}">'
        + _document_defaults(profile.writer_document_defaults)
        + styles
        + "</w:styles>"
    )


def _base_package(profile: StyleProfile) -> bytes:
    """An empty document whose styles are `profile`'s, with Word's current compatibility mode."""
    members = {
        "[Content_Types].xml": _CONTENT_TYPES,
        "_rels/.rels": _PACKAGE_RELATIONSHIPS,
        "word/document.xml": _DECLARATION + f'<w:document xmlns:w="{_W_NAMESPACE}"><w:body/></w:document>',
        "word/_rels/document.xml.rels": _DOCUMENT_RELATIONSHIPS,
        "word/styles.xml": _styles_xml(profile),
        "word/settings.xml": _DECLARATION
        + f'<w:settings xmlns:w="{_W_NAMESPACE}">{_WORD_COMPATIBILITY}</w:settings>',
        "docProps/core.xml": _EMPTY_CORE_PROPERTIES,
    }
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as package:
        for name, text in members.items():
            package.writestr(name, text.encode("utf-8"))
    return buffer.getvalue()


def _stamped(package: bytes, moment: datetime) -> bytes:
    """`package` with every member stamped `moment` (UTC), so the same input is the same bytes."""
    stamp = max(
        (moment.year, moment.month, moment.day, moment.hour, moment.minute, moment.second),
        (1980, 1, 1, 0, 0, 0),
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(package)) as source, zipfile.ZipFile(buffer, "w") as target:
        for member in source.infolist():
            stamped = zipfile.ZipInfo(member.filename, date_time=stamp)
            stamped.compress_type = zipfile.ZIP_DEFLATED
            stamped.external_attr = member.external_attr
            target.writestr(stamped, source.read(member))
    return buffer.getvalue()
