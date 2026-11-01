"""`VerifyManifest` on an exported .docx, and the snapshot-reference check on every manifest (`v1-e06-t07`).

Each document is t02's news card exported by ExportService and changed through python-docx
(`tests/fixtures/verify/docx_world.py`); the use case runs over the world it was cut in, so the
snapshot is the one the card names. Expected verdicts are written by hand from the fixture text: "two"
is evidence [18, 21), "project" is evidence 156, and in the body everything after the cut at 108 is
three characters further on.
"""

from __future__ import annotations

import io
import json
import zipfile
from collections.abc import Callable
from datetime import timedelta
from typing import Any, Final

import docx
import pytest
from docx.document import Document
from tests.fixtures.card.card_world import CardWorld
from tests.fixtures.verification.verification_world import run
from tests.fixtures.verify.docx_world import (
    CRITICS_SENTENCE,
    NEWS_CARD_ID,
    NEWS_CITE,
    NEWS_TAG,
    PROFILES,
    NewsExport,
    edited,
    evidence_paragraph,
    news_card_with_the_critics_sentence_unmarked,
    news_export,
    render_with,
    replace_in_evidence,
    without_markup,
)

from debate_core.application.verify_manifest import (
    CardVerdict,
    DocumentCheck,
    InvalidManifest,
    ManifestVerificationReport,
    NoManifest,
    ReportedRange,
    ReportedReason,
    VerifyManifest,
)
from debate_core.domain import VerificationStatus
from debate_core.evidence.format_profile import load_format_profile, load_format_profiles
from debate_core.evidence.manifest import CardManifest, SnapshotReference, serialize_manifest
from debate_core.evidence.verification_types import ReasonCode, WarningCode

PROFILE_SET: Final = load_format_profiles(PROFILES)


@pytest.fixture(scope="module")
def export() -> NewsExport:
    return news_export()


def use_case(world: CardWorld) -> VerifyManifest:
    return VerifyManifest(
        verifier=world.verifier, articles=world.articles, format_profiles=lambda: PROFILE_SET
    )


def verify(
    export: NewsExport, document: bytes, *, world: CardWorld | None = None
) -> ManifestVerificationReport:
    return run(use_case(world or export.world).verify(document))


def verdict_of(report: ManifestVerificationReport) -> CardVerdict:
    [verdict] = report.cards
    return verdict


# --------------------------------------------------------------------------------------------
# ac1: an untouched export verifies, every card with its snapshot id
# --------------------------------------------------------------------------------------------


def test_docx_an_untouched_export_lists_its_card_verified_with_its_snapshot(export: NewsExport) -> None:
    report = verify(export, export.document)

    verdict = verdict_of(report)
    assert (verdict.card_id, verdict.tag, verdict.status, verdict.reasons, verdict.warnings) == (
        NEWS_CARD_ID,
        NEWS_TAG,
        VerificationStatus.VERIFIED,
        (),
        (),
    )
    assert verdict.snapshot_id == export.snapshot.snapshot.snapshot_id
    assert report.all_verified
    assert report.document == DocumentCheck(
        format_profile="team-default", body_cards=1, unclaimed_body_cards=()
    )


def test_docx_an_export_in_another_profile_is_read_in_the_profile_its_card_names() -> None:
    flow = load_format_profile("flow", PROFILES)
    exported = news_export(profile=flow, change_card=lambda card: card.evolve(format_profile="flow"))

    report = verify(exported, exported.document)

    assert verdict_of(report).status is VerificationStatus.VERIFIED
    assert report.document is not None and report.document.format_profile == "flow"


# --------------------------------------------------------------------------------------------
# ac2: a changed word is MANUAL_EDIT, with the card and the character range
# --------------------------------------------------------------------------------------------


def test_docx_a_changed_word_is_manual_edit_with_the_card_and_its_range(export: NewsExport) -> None:
    report = verify(export, edited(export.document, replace_in_evidence("two years", "three years")))

    verdict = verdict_of(report)
    assert (verdict.position, verdict.status, verdict.reason_codes) == (
        0,
        VerificationStatus.UNVERIFIED,
        (ReasonCode.MANUAL_EDIT,),
    )
    [reason] = verdict.reasons
    assert (reason.evidence_range, reason.body_range) == (
        ReportedRange(start=18, end=21),
        ReportedRange(start=18, end=23),
    )
    assert "evidence_text [18, 21)" in reason.detail
    assert "three" not in reason.detail and "two" not in reason.detail
    assert not report.all_verified


def test_docx_a_word_typed_into_the_evidence_is_manual_edit(export: NewsExport) -> None:
    """ "not" before "realistic" (evidence 192, body 195): an insertion, not a deletion."""
    report = verify(
        export, edited(export.document, replace_in_evidence("never realistic", "never not realistic"))
    )

    [reason] = verdict_of(report).reasons
    assert (reason.code, reason.evidence_range, reason.body_range) == (
        ReasonCode.MANUAL_EDIT,
        ReportedRange(start=192, end=192),
        ReportedRange(start=195, end=199),
    )


# --------------------------------------------------------------------------------------------
# ac3: markup removed or an unmarked sentence deleted stays VERIFIED; an interpolation warns
# --------------------------------------------------------------------------------------------


def test_docx_taking_the_underlining_and_highlighting_off_keeps_the_card_verified(export: NewsExport) -> None:
    verdict = verdict_of(verify(export, edited(export.document, without_markup)))

    assert (verdict.status, verdict.reasons, verdict.warnings) == (VerificationStatus.VERIFIED, (), ())


def test_docx_deleting_an_entire_unmarked_sentence_keeps_the_card_verified() -> None:
    exported = news_export(change_card=news_card_with_the_critics_sentence_unmarked)

    verdict = verdict_of(
        verify(exported, edited(exported.document, replace_in_evidence(CRITICS_SENTENCE, "")))
    )

    assert (verdict.status, verdict.reasons, verdict.warnings) == (VerificationStatus.VERIFIED, (), ())


def test_docx_a_bracketed_interpolation_is_a_warning_on_a_verified_card(export: NewsExport) -> None:
    report = verify(
        export, edited(export.document, replace_in_evidence("the project", "the [mammoth] project"))
    )

    verdict = verdict_of(report)
    assert (verdict.status, verdict.reasons) == (VerificationStatus.VERIFIED, ())
    [warning] = verdict.warnings
    assert (warning.code, warning.evidence_offset) == (WarningCode.INTERPOLATION_PRESENT, 156)
    assert "7 character(s)" in warning.detail and "mammoth" not in warning.detail
    assert report.all_verified


# --------------------------------------------------------------------------------------------
# ac4: no manifest is NO_MANIFEST; a snapshot missing here is SNAPSHOT_MISSING
# --------------------------------------------------------------------------------------------


def hand_cut_team_file() -> bytes:
    """A team file typed in Word: the same tag, cite and evidence, and no manifest."""
    document = docx.Document()
    document.add_paragraph(NEWS_TAG, style="Heading 4")
    document.add_paragraph(NEWS_CITE.full_cite)
    document.add_paragraph("The delay follows two years of failed attempts.")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_docx_a_document_with_no_embedded_manifest_is_no_manifest(export: NewsExport) -> None:
    with pytest.raises(NoManifest, match="verify reads only .docx files debate-research exported"):
        verify(export, hand_cut_team_file())


def test_docx_a_card_whose_snapshot_is_not_stored_here_is_snapshot_missing(export: NewsExport) -> None:
    """The world the card was not cut in holds no snapshot: the verifier's own code, and nothing more."""
    verdict = verdict_of(verify(export, export.document, world=CardWorld()))

    assert (verdict.status, verdict.reason_codes) == (
        VerificationStatus.UNVERIFIED,
        (ReasonCode.SNAPSHOT_MISSING,),
    )


# --------------------------------------------------------------------------------------------
# ac5: a tampered snapshot reference fails with its own code; an untampered one passes
# --------------------------------------------------------------------------------------------

TAMPERINGS: Final[dict[str, Callable[[SnapshotReference], Any]]] = {
    "sha256": lambda reference: "0" * 64,
    "normalized_text_hash": lambda reference: "f" * 64,
    "normalizer_version": lambda reference: "normalizer-v0",
    "extractor_version": lambda reference: "fixture-extractor/2",
    "retrieved_at": lambda reference: reference.retrieved_at + timedelta(seconds=1),
    "canonical_url": lambda reference: "https://www.dailymeridian.example/science/another-story",
}


def tampered(export: NewsExport, field: str) -> CardManifest:
    [reference] = export.manifest.snapshots
    changed = reference.model_copy(update={field: TAMPERINGS[field](reference)})
    return export.manifest.model_copy(update={"snapshots": (changed,)})


@pytest.mark.parametrize("field", list(TAMPERINGS))
@pytest.mark.parametrize("form", ["json", "docx"])
def test_docx_and_json_a_tampered_snapshot_reference_is_its_own_reason_naming_the_field(
    export: NewsExport, field: str, form: str
) -> None:
    manifest = tampered(export, field)
    document = serialize_manifest(manifest) if form == "json" else render_with(manifest, export)

    verdict = verdict_of(verify(export, document))

    assert (verdict.status, verdict.reason_codes) == (
        VerificationStatus.UNVERIFIED,
        (ReasonCode.SNAPSHOT_REFERENCE_MISMATCH,),
    )
    assert verdict.reasons[0].detail.endswith(f"differs from the stored snapshot record in {field}")


@pytest.mark.parametrize("form", ["json", "docx"])
def test_docx_and_json_an_untampered_snapshot_reference_passes(export: NewsExport, form: str) -> None:
    document = serialize_manifest(export.manifest) if form == "json" else export.document

    assert verdict_of(verify(export, document)).status is VerificationStatus.VERIFIED


def test_json_a_version_2_manifest_that_references_another_snapshot_is_a_reference_mismatch(
    export: NewsExport,
) -> None:
    data = json.loads(serialize_manifest(export.manifest))
    data["snapshots"][0]["snapshot_id"] = "0SNAP000000000000000000099"
    document = json.dumps(data).encode()

    verdict = verdict_of(verify(export, document))

    assert verdict.reason_codes == (ReasonCode.SNAPSHOT_REFERENCE_MISMATCH,)
    assert "lists no reference to snapshot" in verdict.reasons[0].detail


# --------------------------------------------------------------------------------------------
# The body as a whole: a card taken out, a card pasted in, a body that cannot be read
# --------------------------------------------------------------------------------------------


def without_the_tag(document: Document) -> None:
    tag = document.paragraphs[0]._p  # pyright: ignore[reportPrivateUsage]
    parent = tag.getparent()
    assert parent is not None
    parent.remove(tag)


def test_docx_a_card_whose_tag_was_taken_out_of_the_body_is_body_card_missing(export: NewsExport) -> None:
    report = verify(export, edited(export.document, without_the_tag))

    assert verdict_of(report).reason_codes == (ReasonCode.BODY_CARD_MISSING,)
    assert report.document == DocumentCheck(
        format_profile="team-default", body_cards=0, unclaimed_body_cards=()
    )


def with_a_pasted_card(document: Document) -> None:
    document.add_paragraph("A card pasted in by hand", style="Heading 4")
    document.add_paragraph("Someone 24")
    document.add_paragraph("Words no snapshot holds.")


def test_docx_a_card_pasted_into_the_body_is_listed_and_changes_no_verdict(export: NewsExport) -> None:
    report = verify(export, edited(export.document, with_a_pasted_card))

    assert verdict_of(report).status is VerificationStatus.VERIFIED
    assert report.document == DocumentCheck(
        format_profile="team-default", body_cards=2, unclaimed_body_cards=(1,)
    )


def test_docx_a_paragraph_added_under_the_evidence_is_manual_edit(export: NewsExport) -> None:
    def added(document: Document) -> None:
        assert evidence_paragraph(document).text, "the evidence is the last paragraph before this one"
        document.add_paragraph("And the revival will never work.")

    assert verdict_of(verify(export, edited(export.document, added))).reason_codes == (
        ReasonCode.MANUAL_EDIT,
    )


def test_docx_cards_naming_a_profile_this_build_lacks_cannot_be_read() -> None:
    exported = news_export(change_card=lambda card: card.evolve(format_profile="camp-style"))

    with pytest.raises(InvalidManifest) as refused:
        verify(exported, exported.document)

    assert refused.value.json_path == "$.cards[0].format_profile"
    assert "does not have" in refused.value.problem


def test_docx_a_body_that_cannot_be_read_is_an_invalid_document(export: NewsExport) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(export.document)) as source, zipfile.ZipFile(buffer, "w") as target:
        for member in source.infolist():
            content = source.read(member)
            if member.filename == "word/document.xml":
                content = content.replace(b"?>", b"?><!DOCTYPE w:document>", 1)
            target.writestr(member, content)

    with pytest.raises(InvalidManifest, match="body cannot be read"):
        verify(export, buffer.getvalue())


def test_docx_a_manifest_part_that_is_not_xml_is_an_invalid_manifest(export: NewsExport) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(export.document)) as source, zipfile.ZipFile(buffer, "w") as target:
        for member in source.infolist():
            content = source.read(member)
            if member.filename == "customXml/item1.xml":
                content = content[: len(content) // 2]
            target.writestr(member, content)

    with pytest.raises(InvalidManifest, match="embedded manifest cannot be read"):
        verify(export, buffer.getvalue())


def test_docx_the_manifest_is_checked_before_the_body(export: NewsExport) -> None:
    """A tampered card in the manifest is the verifier's TEXT_MISMATCH even though the body matches it."""
    [card] = export.manifest.cards
    changed = card.evolve(evidence_text=card.evidence_text.replace("two", "six", 1))
    manifest = export.manifest.model_copy(update={"cards": (changed,)})

    verdict = verdict_of(verify(export, render_with(manifest, export)))

    assert verdict.reason_codes == (ReasonCode.TEXT_MISMATCH,)


def test_docx_a_verdict_cannot_be_verified_and_carry_a_reason(export: NewsExport) -> None:
    """Defence in depth: whatever finds a reason beside the verifier, the verdict cannot stay VERIFIED."""
    verdict = verdict_of(verify(export, export.document))
    manual_edit = ReportedReason(code=ReasonCode.MANUAL_EDIT, detail="the body changes evidence_text [0, 1)")

    with pytest.raises(ValueError, match="VERIFIED exactly when it carries no reason"):
        CardVerdict.model_validate({**verdict.model_dump(exclude={"reason_codes"}), "reasons": [manual_edit]})
    with pytest.raises(ValueError, match="VERIFIED exactly when it carries no reason"):
        CardVerdict.model_validate({**verdict.model_dump(exclude={"reason_codes"}), "status": "UNVERIFIED"})
