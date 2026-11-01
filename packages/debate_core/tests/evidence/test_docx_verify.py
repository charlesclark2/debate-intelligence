"""Tests for reading a rendered .docx's body and comparing it with its cards (`v1-e06-t07`).

Two plan nodes, by name: tests with ``reader`` in their name are the body reader's
(:mod:`debate_core.evidence.docx_reader`), tests with ``diff`` the comparison's
(:mod:`debate_core.evidence.docx_verify`). Documents are real exports of t02's news card
(`tests/fixtures/verify/docx_world.py`), changed through python-docx the way Word changes them. Every
expected text and offset is counted by hand from the fixture's evidence, whose pieces the world names:
the first sentence is evidence [0, 108), the cut sits at 108, the critics' sentence is [108, 202) and
the rest [202, 398). In an untouched body everything after the cut is three characters further on,
the marker's length.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable
from typing import Final

import pytest
from docx.document import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from tests.fixtures.verify.docx_world import (
    CRITICS_SENTENCE,
    FIRST_SENTENCE,
    LAST_PART,
    MARKER,
    NEWS_CITE,
    NEWS_TAG,
    RENDERED_EVIDENCE,
    TEAM_DEFAULT,
    NewsExport,
    edited,
    evidence_paragraph,
    news_export,
)

from debate_core.domain import Card, Interpolation, VerificationStatus
from debate_core.evidence.card_rendering import CardToRender, DocumentHeadings
from debate_core.evidence.docx_reader import UnreadableDocument, read_document_body
from debate_core.evidence.docx_renderer import render_cards_docx
from debate_core.evidence.docx_verify import (
    BodyCard,
    BodyInterpolation,
    BodyParagraph,
    CharacterRange,
    DocumentBody,
    ManualEdit,
    body_evidence,
    compare_body,
    match_body_cards,
)
from debate_core.evidence.manifest import CardManifest

STYLES: Final = TEAM_DEFAULT.style_profile
OBJECT: Final = "￼"
"""What the reader reads in place of anything Word shows that is not text: a picture, a symbol."""
SHIFT: Final = len(MARKER)


@pytest.fixture(scope="module")
def export() -> NewsExport:
    return news_export()


@pytest.fixture(scope="module")
def card(export: NewsExport) -> Card:
    return export.card


def body_of(document: bytes) -> DocumentBody:
    return read_document_body(document, STYLES)


def with_paragraph_xml(xml: str) -> Callable[[Document], None]:
    """Append ``xml`` (one ``w:p`` or ``w:tbl``, ``w`` declared by the caller) after the evidence."""

    def change(document: Document) -> None:
        evidence_paragraph(document)._p.addnext(parse_xml(xml))  # pyright: ignore[reportPrivateUsage]

    return change


def with_runs_in_evidence(xml: str) -> Callable[[Document], None]:
    """Insert run-level ``xml`` into the evidence paragraph right after its first run."""

    def change(document: Document) -> None:
        first = evidence_paragraph(document).runs[0]._r  # pyright: ignore[reportPrivateUsage]
        first.addnext(parse_xml(xml))

    return change


# --------------------------------------------------------------------------------------------
# Node docx-reader: the body as Word shows it, cut into cards by the profile's styles
# --------------------------------------------------------------------------------------------


def test_reader_reads_an_untouched_export_as_its_tag_its_cite_and_its_evidence(export: NewsExport) -> None:
    assert body_of(export.document) == DocumentBody(
        cards=(
            BodyCard(
                position=0,
                tag=NEWS_TAG,
                paragraphs=(
                    BodyParagraph(NEWS_CITE.full_cite, is_cite=True),
                    BodyParagraph(RENDERED_EVIDENCE),
                ),
            ),
        )
    )


def test_reader_leaves_headings_out_and_starts_a_card_at_each_tag(export: NewsExport) -> None:
    second = export.card.evolve(card_id="0CVT0000000000000000000002", tag="A second tag")
    manifest = CardManifest(
        manifest_version=2,
        generated_at=export.manifest.generated_at,
        cards=(export.card, second),
        snapshots=export.manifest.snapshots,
    )
    document = render_cards_docx(
        manifest,
        [CardToRender(export.card, NEWS_CITE), CardToRender(second, NEWS_CITE)],
        TEAM_DEFAULT,
        headings=DocumentHeadings(pocket="1AC", hat="De-extinction", block="Funding"),
    )

    body = body_of(document)

    assert [(card.position, card.tag, len(card.paragraphs)) for card in body.cards] == [
        (0, NEWS_TAG, 2),
        (1, "A second tag", 2),
    ]


def test_reader_reads_the_banner_of_an_unverified_card_as_a_paragraph_before_the_cite(
    export: NewsExport,
) -> None:
    unverified = export.card.evolve(verification_status=VerificationStatus.UNVERIFIED)
    manifest = export.manifest.model_copy(update={"cards": (unverified,)})
    document = render_cards_docx(
        manifest, [CardToRender(unverified, NEWS_CITE, ("CITATION_UNVERIFIED",))], TEAM_DEFAULT
    )

    [card] = body_of(document).cards

    assert card.paragraphs == (
        BodyParagraph("UNVERIFIED — not a finished card: CITATION_UNVERIFIED"),
        BodyParagraph(NEWS_CITE.full_cite, is_cite=True),
        BodyParagraph(RENDERED_EVIDENCE),
    )


def test_reader_ends_a_card_at_a_heading_of_any_level(export: NewsExport) -> None:
    heading = (
        f'<w:p {nsdecls("w")}><w:pPr><w:pStyle w:val="Heading1"/></w:pPr>'
        "<w:r><w:t>Next pocket</w:t></w:r></w:p>"
    )
    after = f"<w:p {nsdecls('w')}><w:r><w:t>Not under the card</w:t></w:r></w:p>"

    def change(document: Document) -> None:
        evidence = evidence_paragraph(document)._p  # pyright: ignore[reportPrivateUsage]
        evidence.addnext(parse_xml(after))
        evidence.addnext(parse_xml(heading))

    [card] = body_of(edited(export.document, change)).cards

    assert [paragraph.text for paragraph in card.paragraphs] == [NEWS_CITE.full_cite, RENDERED_EVIDENCE]


def test_reader_reads_a_tracked_insertion_and_leaves_out_a_tracked_deletion(export: NewsExport) -> None:
    """Word shows an insertion and, in the final view, not a deletion; the reader reads what it shows."""
    runs = (
        f'<w:ins {nsdecls("w")} w:id="1" w:author="s" w:date="2026-10-30T00:00:00Z">'
        '<w:r><w:t xml:space="preserve">not </w:t></w:r></w:ins>'
    )
    deletion = (
        f'<w:del {nsdecls("w")} w:id="2" w:author="s" w:date="2026-10-30T00:00:00Z"><w:r>'
        "<w:delText>gone</w:delText></w:r></w:del>"
    )

    document = edited(export.document, with_runs_in_evidence(runs))
    document = edited(document, with_runs_in_evidence(deletion))

    evidence = body_of(document).cards[0].paragraphs[-1].text
    assert evidence == "The delay follows two years of " + "not " + RENDERED_EVIDENCE[31:]


def test_reader_leaves_out_hidden_text(export: NewsExport) -> None:
    hidden = f"<w:r {nsdecls('w')}><w:rPr><w:vanish/></w:rPr><w:t>hidden</w:t></w:r>"

    evidence = body_of(edited(export.document, with_runs_in_evidence(hidden))).cards[0].paragraphs[-1].text

    assert evidence == RENDERED_EVIDENCE


@pytest.mark.parametrize(
    "run_xml",
    [
        f'<w:r {nsdecls("w")}><w:sym w:font="Symbol" w:char="F0AE"/></w:r>',
        f'<w:r {nsdecls("w")}><w:footnoteReference w:id="1"/></w:r>',
        f"<w:r {nsdecls('w', 'wp')}><w:drawing><wp:inline/></w:drawing></w:r>",
    ],
    ids=["a symbol", "a footnote mark", "a picture"],
)
def test_reader_reads_what_word_shows_that_is_not_text_as_an_object_character(
    export: NewsExport, run_xml: str
) -> None:
    evidence = body_of(edited(export.document, with_runs_in_evidence(run_xml))).cards[0].paragraphs[-1].text

    assert evidence == RENDERED_EVIDENCE[:31] + OBJECT + RENDERED_EVIDENCE[31:]


def test_reader_reads_tabs_breaks_and_text_inside_a_hyperlink(export: NewsExport) -> None:
    runs = (
        f'<w:hyperlink {nsdecls("w")} w:anchor="x"><w:r><w:tab/><w:t>linked</w:t><w:br/></w:r></w:hyperlink>'
    )

    evidence = body_of(edited(export.document, with_runs_in_evidence(runs))).cards[0].paragraphs[-1].text

    assert evidence == RENDERED_EVIDENCE[:31] + "\tlinked\n" + RENDERED_EVIDENCE[31:]


def test_reader_reads_a_table_under_a_card_as_a_paragraph_of_its_text(export: NewsExport) -> None:
    table = (
        f"<w:tbl {nsdecls('w')}><w:tr><w:tc><w:p><w:r><w:t>left</w:t></w:r></w:p></w:tc>"
        "<w:tc><w:p><w:r><w:t>right</w:t></w:r></w:p></w:tc></w:tr></w:tbl>"
    )

    [card] = body_of(edited(export.document, with_paragraph_xml(table))).cards

    assert card.paragraphs[-1] == BodyParagraph("left\nright")


def test_reader_knows_a_tag_by_its_style_name_when_word_has_renamed_its_id(export: NewsExport) -> None:
    """German Word writes heading 4's id as `berschrift4`; the name stays `heading 4`."""

    def localized(document: Document) -> None:
        document.styles["heading 4"].element.set(qn("w:styleId"), "berschrift4")
        document.paragraphs[0]._p.pPr.pStyle.set(qn("w:val"), "berschrift4")  # pyright: ignore[reportOptionalMemberAccess, reportPrivateUsage]

    [card] = body_of(edited(export.document, localized)).cards

    assert card.tag == NEWS_TAG


def test_reader_refuses_a_file_that_is_not_a_docx() -> None:
    with pytest.raises(UnreadableDocument, match="not a .docx"):
        body_of(b"Marrow basin wells, card list\n")


def test_reader_refuses_a_package_with_no_main_document(export: NewsExport) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(export.document)) as source, zipfile.ZipFile(buffer, "w") as target:
        for member in source.infolist():
            if member.filename != "word/document.xml":
                target.writestr(member, source.read(member))

    with pytest.raises(UnreadableDocument, match="main document"):
        body_of(buffer.getvalue())


def test_reader_refuses_a_main_document_that_declares_a_document_type(export: NewsExport) -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(export.document)) as source, zipfile.ZipFile(buffer, "w") as target:
        for member in source.infolist():
            content = source.read(member)
            if member.filename == "word/document.xml":
                content = content.replace(b"?>", b'?><!DOCTYPE w:document [<!ENTITY e "x">]>', 1)
            target.writestr(member, content)

    with pytest.raises(UnreadableDocument, match="document type"):
        body_of(buffer.getvalue())


# --------------------------------------------------------------------------------------------
# Node body-diff: which body card is which, which paragraphs are evidence, what changed
# --------------------------------------------------------------------------------------------


def body_card(*paragraphs: BodyParagraph, position: int = 0, tag: str = NEWS_TAG) -> BodyCard:
    return BodyCard(position=position, tag=tag, paragraphs=paragraphs)


CITE: Final = BodyParagraph(NEWS_CITE.full_cite, is_cite=True)


def test_diff_matches_body_cards_by_position_when_the_counts_agree_whatever_their_tags() -> None:
    body = DocumentBody(
        cards=(body_card(tag="Rewritten by the student"), body_card(position=1, tag="Second"))
    )

    matched, unclaimed = match_body_cards(["First", "Second"], body)

    assert matched == body.cards
    assert unclaimed == ()


def test_diff_matches_by_tag_in_order_when_a_card_was_taken_out_of_the_body() -> None:
    kept = body_card(tag="Second")
    matched, unclaimed = match_body_cards(["First", "Second"], DocumentBody(cards=(kept,)))

    assert matched == (None, kept)
    assert unclaimed == ()


def test_diff_reports_a_body_card_no_manifest_card_claims() -> None:
    ours, theirs = body_card(tag="First"), body_card(position=1, tag="Pasted in by hand")
    body = DocumentBody(cards=(ours, body_card(position=2, tag="Also pasted"), theirs))

    matched, unclaimed = match_body_cards(["First", "Pasted in by hand"], body)

    assert matched == (ours, theirs)
    assert [card.tag for card in unclaimed] == ["Also pasted"]


def test_diff_takes_the_evidence_from_after_the_cite_and_leaves_blank_paragraphs_out(card: Card) -> None:
    shown = body_card(
        CITE, BodyParagraph("first"), BodyParagraph("  "), BodyParagraph(""), BodyParagraph("second")
    )

    assert body_evidence(card, shown) == "first\nsecond"


def test_diff_finds_the_cite_by_place_when_no_paragraph_carries_the_cite_style(card: Card) -> None:
    plain = BodyParagraph(NEWS_CITE.full_cite)

    assert body_evidence(card, body_card(plain, BodyParagraph("evidence"))) == "evidence"
    unverified = card.evolve(verification_status=VerificationStatus.UNVERIFIED)
    assert body_evidence(
        unverified, body_card(BodyParagraph("banner"), plain, BodyParagraph("evidence"))
    ) == ("evidence")


def test_diff_an_untouched_body_follows_the_policy_with_nothing_removed(card: Card) -> None:
    comparison = compare_body(card, RENDERED_EVIDENCE, MARKER)

    assert (comparison.manual_edits, comparison.interpolations, comparison.deletions) == ((), (), ())


def test_diff_a_changed_word_is_a_manual_edit_with_both_ranges(card: Card) -> None:
    """ "two" is evidence [18, 21); "three" is body [18, 23)."""
    body = RENDERED_EVIDENCE.replace("two years", "three years", 1)

    assert compare_body(card, body, MARKER).manual_edits == (
        ManualEdit(evidence_range=CharacterRange(18, 21), body_range=CharacterRange(18, 23)),
    )


def test_diff_a_word_whose_case_changed_is_a_manual_edit(card: Card) -> None:
    """ "Every" is evidence [204, 209) and body [207, 212): a capital is part of the quotation."""
    body = RENDERED_EVIDENCE.replace("“Every", "“every", 1)

    assert compare_body(card, body, MARKER).manual_edits == (
        ManualEdit(evidence_range=CharacterRange(204, 209), body_range=CharacterRange(207, 212)),
    )


def test_diff_letters_taken_out_of_a_word_are_a_manual_edit_not_a_deletion(card: Card) -> None:
    """ "failed" is evidence [31, 37); "fail" is body [31, 35)."""
    body = RENDERED_EVIDENCE.replace("failed", "fail", 1)

    assert compare_body(card, body, MARKER).manual_edits == (
        ManualEdit(evidence_range=CharacterRange(31, 37), body_range=CharacterRange(31, 35)),
    )


def test_diff_deleting_the_space_between_two_words_is_a_manual_edit(card: Card) -> None:
    """ "two years" [18, 27) becomes one word, "twoyears", body [18, 26)."""
    body = RENDERED_EVIDENCE.replace("two years", "twoyears", 1)

    assert compare_body(card, body, MARKER).manual_edits == (
        ManualEdit(evidence_range=CharacterRange(18, 27), body_range=CharacterRange(18, 26)),
    )


@pytest.mark.parametrize(
    "body",
    [
        FIRST_SENTENCE + MARKER + LAST_PART,
        FIRST_SENTENCE + MARKER.rstrip() + LAST_PART,
        FIRST_SENTENCE + LAST_PART,
    ],
    ids=["marker kept, its space doubled", "marker kept", "marker deleted too"],
)
def test_diff_deleting_a_whole_sentence_is_a_deletion_that_follows_the_policy(card: Card, body: str) -> None:
    comparison = compare_body(card, body, MARKER)

    assert comparison.manual_edits == ()
    assert comparison.deletions == (CharacterRange(108, 202),)


def test_diff_deleting_words_and_typing_the_marker_in_their_place_follows_the_policy(card: Card) -> None:
    """ "failed attempts to grow edited" is evidence [31, 61)."""
    body = RENDERED_EVIDENCE.replace("failed attempts to grow edited", "…", 1)

    comparison = compare_body(card, body, MARKER)

    assert comparison.manual_edits == ()
    assert comparison.deletions == (CharacterRange(31, 61),)


def test_diff_a_marker_where_nothing_was_cut_is_a_manual_edit(card: Card) -> None:
    """After "realistic." (evidence 202, body 205), where nothing was cut."""
    body = RENDERED_EVIDENCE.replace("realistic.", "realistic.…", 1)

    assert compare_body(card, body, MARKER).manual_edits == (
        ManualEdit(evidence_range=CharacterRange(202, 202), body_range=CharacterRange(205, 206)),
    )


def test_diff_the_marker_is_the_profiles_own(card: Card) -> None:
    """A profile whose marker is three full stops: one typed where words were cut follows the policy."""
    body = FIRST_SENTENCE + " ... " + CRITICS_SENTENCE + LAST_PART

    assert compare_body(card, body, " ... ").manual_edits == ()
    assert compare_body(card, body, MARKER).manual_edits != ()


def test_diff_an_added_interpolation_is_reported_where_it_is_read_and_follows_the_policy(card: Card) -> None:
    """ "project" is evidence 156; the interpolation is read before it."""
    body = RENDERED_EVIDENCE.replace("the project", "the [mammoth] project", 1)

    comparison = compare_body(card, body, MARKER)

    assert comparison.manual_edits == ()
    assert comparison.interpolations == (BodyInterpolation(anchor=156, length=7, recorded=False),)


def test_diff_an_interpolation_the_card_records_is_reported_as_recorded(card: Card) -> None:
    recorded = card.evolve(interpolations=(Interpolation(anchor=156, text="mammoth revival"),))
    body = (
        FIRST_SENTENCE
        + MARKER
        + CRITICS_SENTENCE.replace("project", "[mammoth revival]project", 1)
        + LAST_PART
    )

    comparison = compare_body(recorded, body, MARKER)

    assert comparison.manual_edits == ()
    assert comparison.interpolations == (BodyInterpolation(anchor=156, length=15, recorded=True),)


def test_diff_words_replaced_by_an_interpolation_follow_the_policy(card: Card) -> None:
    """Deleting "Conservation biologists" [108, 131) and bracketing words of one's own in their place."""
    body = RENDERED_EVIDENCE.replace("Conservation biologists", "[Scientists]", 1)

    comparison = compare_body(card, body, MARKER)

    assert comparison.manual_edits == ()
    assert comparison.deletions == (CharacterRange(108, 131),)
    assert [item.length for item in comparison.interpolations] == [10]


@pytest.mark.parametrize(
    ("inserted", "body_end"),
    [("[] ", 162), ("[  ] ", 164), ("[a] [b] ", 167)],
    ids=["empty brackets", "brackets around spaces", "two interpolations in one place"],
)
def test_diff_brackets_an_interpolation_cannot_be_are_a_manual_edit(
    card: Card, inserted: str, body_end: int
) -> None:
    body = RENDERED_EVIDENCE.replace("the project", "the " + inserted + "project", 1)

    [edit] = compare_body(card, body, MARKER).manual_edits

    assert (edit.evidence_range, edit.body_range) == (CharacterRange(156, 156), CharacterRange(159, body_end))


def test_diff_an_added_space_is_a_manual_edit(card: Card) -> None:
    """Between "“" (evidence 203) and "Every" (204)."""
    body = RENDERED_EVIDENCE.replace("“Every", "“ Every", 1)

    assert compare_body(card, body, MARKER).manual_edits == (
        ManualEdit(evidence_range=CharacterRange(204, 204), body_range=CharacterRange(207, 208)),
    )


@pytest.mark.parametrize("body", ["", "   ", MARKER, "[the card was here]"])
def test_diff_a_body_that_keeps_none_of_the_evidence_is_one_manual_edit_of_all_of_it(
    card: Card, body: str
) -> None:
    assert compare_body(card, body, MARKER).manual_edits == (
        ManualEdit(evidence_range=CharacterRange(0, 398), body_range=CharacterRange(0, len(body))),
    )


def test_diff_a_paragraph_added_under_the_evidence_is_a_manual_edit(card: Card) -> None:
    shown = body_card(CITE, BodyParagraph(RENDERED_EVIDENCE), BodyParagraph("And it will never work."))

    [edit] = compare_body(card, body_evidence(card, shown), MARKER).manual_edits

    assert (edit.evidence_range, edit.body_range) == (CharacterRange(398, 398), CharacterRange(401, 425))


def test_diff_a_body_too_different_to_compare_is_one_manual_edit_of_all_of_it(card: Card) -> None:
    body = "word " * 1100

    assert compare_body(card, body, MARKER).manual_edits == (
        ManualEdit(evidence_range=CharacterRange(0, 398), body_range=CharacterRange(0, len(body))),
    )


def test_diff_text_the_card_itself_brackets_is_quoted_text_not_an_interpolation(card: Card) -> None:
    """A source may bracket its own words; the same group in the body is the source's, not the student's.

    Only the comparison is under test, so the card's text is swapped without re-cutting it."""
    text = "Shares fell [sharply] on Monday."
    bracketed = card.model_copy(update={"evidence_text": text, "omitted_ranges": ()})

    comparison = compare_body(bracketed, text, MARKER)

    assert (comparison.manual_edits, comparison.interpolations) == ((), ())
